import csv
import json
import random
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
import sys
from typing import Any


PROJECT_DIR = Path(__file__).resolve().parents[2]
RAW_DIR = PROJECT_DIR / "data" / "raw"
PROCESSED_DIR = PROJECT_DIR / "data" / "processed"
REPORTS_DIR = PROJECT_DIR / "reports"

DATA_ORG_PATH = RAW_DIR / "data_org1.csv"
MERGED_QUESTIONS_PATH = RAW_DIR / "merged_questions_label_id.xlsx"
CLASS_PATH = PROCESSED_DIR / "class.txt"
TRAIN_PATH = PROCESSED_DIR / "train.txt"
DEV_PATH = PROCESSED_DIR / "dev.txt"
TEST_PATH = PROCESSED_DIR / "test.txt"
REPORT_PATH = REPORTS_DIR / "data_org_preprocess_report.json"

EXPECTED_CLASS_COUNT = 15
RANDOM_SEED = 20260603
OPTION_MARKER_RE = re.compile(r"(?<![A-Za-z0-9\u4e00-\u9fff])([A-D])[.．、]\s*")
OPTION_CUES = [
    "正确的是",
    "错误的是",
    "不正确的是",
    "合理的是",
    "不合理的是",
    "叙述",
    "判断",
    "推断",
    "分析",
    "顺序是",
]


@dataclass(frozen=True)
class Sample:
    text: str
    label_id: int
    source: str


class PreprocessError(RuntimeError):
    pass


def normalize_space(value: Any) -> str:
    if value is None:
        return ""
    text = str(value)
    text = text.replace("\u3000", " ")
    text = text.replace("\xa0", " ")
    text = re.sub(r"[\r\n\t]+", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def strip_common_prefix(value: Any) -> str:
    text = normalize_space(value)
    return re.sub(r"^(题干|选项|答案|解析|答案解析)[:：\s]*", "", text).strip()


def strip_embedded_options(text: str) -> str:
    matches = list(OPTION_MARKER_RE.finditer(text))
    if not matches:
        return text

    for match_index, match in enumerate(matches):
        prefix = text[: match.start()].rstrip()
        suffix_match_count = len(matches) - match_index
        has_question_cue = any(cue in prefix[-100:] for cue in OPTION_CUES)
        if suffix_match_count >= 2 or has_question_cue:
            return normalize_space(prefix.rstrip(" ，,。；;:：("))
    return text


def build_question_text(value: Any) -> str:
    return strip_embedded_options(strip_common_prefix(value))


def build_text(*parts: tuple[str, Any]) -> str:
    segments = []
    for label, value in parts:
        cleaned = strip_common_prefix(value)
        if cleaned:
            segments.append(f"{label}：{cleaned}")
    return normalize_space(" ".join(segments))


def normalize_for_dedupe(text: str) -> str:
    return re.sub(r"\s+", "", text).lower()


def extract_question_text(text: str) -> str:
    question_text = text
    for marker in [" 选项：", " 答案：", " 解析："]:
        if marker in question_text:
            question_text = question_text.split(marker, 1)[0]
    return question_text


def load_classes() -> list[str]:
    if not CLASS_PATH.exists():
        raise PreprocessError(f"class.txt 不存在: {CLASS_PATH}")
    labels = [line.strip() for line in CLASS_PATH.read_text(encoding="utf-8").splitlines() if line.strip()]
    if len(labels) != EXPECTED_CLASS_COUNT:
        raise PreprocessError(
            f"class.txt 必须包含 {EXPECTED_CLASS_COUNT} 个分类，当前为 {len(labels)} 个: {CLASS_PATH}"
        )
    return labels


def ensure_sources_exist() -> None:
    missing = [str(path) for path in [DATA_ORG_PATH, MERGED_QUESTIONS_PATH] if not path.exists()]
    if missing:
        raise PreprocessError(f"raw 源文件不存在: {', '.join(missing)}")


def normalize_label(raw_value: Any, source: str, label_count: int) -> int | None:
    if raw_value is None or str(raw_value).strip() == "":
        return None
    try:
        label = int(float(str(raw_value).strip()))
    except ValueError:
        return None

    if source == "data_org1":
        if 1 <= label <= label_count:
            return label - 1
        if 0 <= label < label_count:
            return label
        return None

    if source == "merged_questions_label_id":
        if 0 <= label < label_count:
            return label
        if label == label_count:
            return label - 1
        return None

    return None


def infer_label_from_text(text: str, labels: list[str]) -> int | None:
    question_text = extract_question_text(text)
    keyword_map = {
        0: ["孟德尔遗传定律", "分离定律", "自由组合定律", "性状分离", "测交实验"],
        1: ["减数分裂与受精作用", "减数分裂", "受精作用", "精子形成", "卵细胞形成"],
        2: ["伴性遗传", "红绿色盲", "伴X", "伴Y", "X染色体隐性", "X染色体显性"],
        3: ["DNA是主要遗传物质", "肺炎链球菌转化实验", "噬菌体侵染细菌实验", "格里菲思", "艾弗里", "赫尔希", "蔡斯"],
        4: ["DNA结构与复制", "DNA双螺旋结构", "DNA复制", "半保留复制"],
        5: ["基因指导蛋白质合成", "蛋白质合成", "遗传信息的转录", "遗传信息的翻译", "转录", "翻译", "密码子", "反密码子", "mRNA", "tRNA", "rRNA"],
        6: ["基因表达与性状关系", "基因控制性状", "基因与性状的关系", "表观遗传"],
        7: ["基因突变和基因重组", "基因突变", "基因重组", "突变和重组"],
        8: ["染色体变异", "染色体结构变异", "染色体数目变异", "染色体组", "单倍体", "多倍体", "三倍体"],
        9: ["人类遗传病", "遗传病发病率", "遗传病概念", "遗传病调查", "遗传病预防", "单基因遗传病", "多基因遗传病", "产前诊断", "遗传咨询", "禁止近亲结婚"],
        10: ["基因在染色体上", "摩尔根", "萨顿", "果蝇杂交"],
        11: ["现代生物进化理论", "自然选择学说", "基因频率", "种群基因库", "物种形成", "隔离"],
        12: ["生物共同祖先证据", "共同祖先", "化石证据", "比较解剖学", "胚胎学证据"],
        13: ["协同进化与生物多样性", "协同进化", "共同进化", "生物多样性"],
        14: ["探究实践与模型实验", "模拟实验", "模型实验", "调查实验"],
    }

    exact_matches = {label_id for label_id, label in enumerate(labels) if label and label in question_text}
    if len(exact_matches) == 1:
        return next(iter(exact_matches))

    matches = set()
    for label_id, keywords in keyword_map.items():
        if any(keyword in question_text for keyword in keywords):
            matches.add(label_id)
    if len(matches) == 1:
        return next(iter(matches))

    return None


def require_fields(actual_fields: list[str], required_fields: list[str], source: str) -> None:
    missing = [field for field in required_fields if field not in actual_fields]
    if missing:
        raise PreprocessError(f"{source} 缺少必需字段: {', '.join(missing)}")


def correct_label(label_id: int, text: str, labels: list[str], report: dict) -> int:
    inferred = infer_label_from_text(text, labels)
    if inferred is None or inferred == label_id:
        return label_id
    report["label_corrections"] += 1
    report["label_correction_examples"].append(
        {
            "from": label_id,
            "to": inferred,
            "text": text[:160],
        }
    )
    return inferred


def load_data_org(labels: list[str], report: dict) -> list[Sample]:
    required_fields = ["question", "knowledge_points_id"]
    samples = []
    source_report = report["sources"]["data_org1"]

    with DATA_ORG_PATH.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        fieldnames = reader.fieldnames or []
        source_report["fields"] = fieldnames
        require_fields(fieldnames, required_fields, "data_org1.csv")
        for row in reader:
            source_report["read_rows"] += 1
            label_id = normalize_label(row.get("knowledge_points_id"), "data_org1", len(labels))
            if label_id is None:
                report["filtered"]["invalid_label"] += 1
                continue
            text = build_question_text(row.get("question"))
            if not text:
                report["filtered"]["empty_text"] += 1
                continue
            label_id = correct_label(label_id, text, labels, report)
            samples.append(Sample(text=text, label_id=label_id, source="data_org1"))

    source_report["valid_rows"] = len(samples)
    return samples


def load_merged_questions(labels: list[str], report: dict) -> list[Sample]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise PreprocessError("缺少 openpyxl，无法读取 merged_questions_label_id.xlsx") from error

    required_fields = ["stem", "label_id"]
    samples = []
    source_report = report["sources"]["merged_questions_label_id"]

    workbook = load_workbook(MERGED_QUESTIONS_PATH, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration as error:
        raise PreprocessError("merged_questions_label_id.xlsx 为空") from error

    headers = [normalize_space(value) for value in header_row]
    source_report["fields"] = headers
    require_fields(headers, required_fields, "merged_questions_label_id.xlsx")
    indexes = {field: headers.index(field) for field in required_fields}

    for row in rows:
        source_report["read_rows"] += 1
        label_id = normalize_label(row[indexes["label_id"]], "merged_questions_label_id", len(labels))
        if label_id is None:
            report["filtered"]["invalid_label"] += 1
            continue
        text = build_question_text(row[indexes["stem"]])
        if not text:
            report["filtered"]["empty_text"] += 1
            continue
        label_id = correct_label(label_id, text, labels, report)
        samples.append(Sample(text=text, label_id=label_id, source="merged_questions_label_id"))

    source_report["valid_rows"] = len(samples)
    return samples


def deduplicate(samples: list[Sample], report: dict) -> list[Sample]:
    seen = set()
    unique = []
    for sample in samples:
        key = (normalize_for_dedupe(sample.text), sample.label_id)
        if key in seen:
            report["filtered"]["duplicate"] += 1
            continue
        seen.add(key)
        unique.append(sample)
    return unique


def split_samples(samples: list[Sample]) -> dict[str, list[Sample]]:
    grouped: dict[int, list[Sample]] = defaultdict(list)
    for sample in samples:
        grouped[sample.label_id].append(sample)

    rng = random.Random(RANDOM_SEED)
    splits = {"train": [], "dev": [], "test": []}
    for label_id in sorted(grouped):
        label_samples = grouped[label_id]
        rng.shuffle(label_samples)
        total = len(label_samples)
        if total >= 10:
            dev_count = max(1, round(total * 0.1))
            test_count = max(1, round(total * 0.1))
        elif total >= 3:
            dev_count = 1
            test_count = 1
        elif total == 2:
            dev_count = 1
            test_count = 0
        else:
            dev_count = 0
            test_count = 0

        train_count = total - dev_count - test_count
        if train_count <= 0:
            train_count = 1
            if test_count > 0:
                test_count -= 1
            elif dev_count > 0:
                dev_count -= 1

        splits["dev"].extend(label_samples[:dev_count])
        splits["test"].extend(label_samples[dev_count : dev_count + test_count])
        splits["train"].extend(label_samples[dev_count + test_count :])

    for rows in splits.values():
        rng.shuffle(rows)
    return splits


def write_split(path: Path, samples: list[Sample]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        "\n".join(f"{sample.text}\t{sample.label_id}" for sample in samples) + "\n",
        encoding="utf-8",
    )


def label_distribution(samples: list[Sample]) -> dict[str, int]:
    return {str(label): count for label, count in sorted(Counter(sample.label_id for sample in samples).items())}


def build_report(labels: list[str]) -> dict:
    return {
        "class_file": {
            "path": str(CLASS_PATH),
            "count": len(labels),
            "labels": labels,
            "unchanged": None,
        },
        "text_mode": "question_stem_only",
        "sources": {
            "data_org1": {"path": str(DATA_ORG_PATH), "fields": [], "read_rows": 0, "valid_rows": 0},
            "merged_questions_label_id": {
                "path": str(MERGED_QUESTIONS_PATH),
                "fields": [],
                "read_rows": 0,
                "valid_rows": 0,
            },
        },
        "filtered": {"invalid_label": 0, "empty_text": 0, "duplicate": 0},
        "label_corrections": 0,
        "label_correction_examples": [],
        "distribution": {},
        "splits": {},
        "outputs": {
            "train": str(TRAIN_PATH),
            "dev": str(DEV_PATH),
            "test": str(TEST_PATH),
            "class": str(CLASS_PATH),
        },
    }


def main() -> int:
    try:
        ensure_sources_exist()
        class_before = CLASS_PATH.read_text(encoding="utf-8") if CLASS_PATH.exists() else ""
        labels = load_classes()
        report = build_report(labels)

        samples = []
        samples.extend(load_data_org(labels, report))
        samples.extend(load_merged_questions(labels, report))
        unique_samples = deduplicate(samples, report)
        if not unique_samples:
            raise PreprocessError("清洗后没有可用样本")

        report["distribution"]["all"] = label_distribution(unique_samples)
        splits = split_samples(unique_samples)
        write_split(TRAIN_PATH, splits["train"])
        write_split(DEV_PATH, splits["dev"])
        write_split(TEST_PATH, splits["test"])

        class_after = CLASS_PATH.read_text(encoding="utf-8")
        report["class_file"]["unchanged"] = class_before == class_after
        if not report["class_file"]["unchanged"]:
            raise PreprocessError("data/processed/class.txt 在清洗过程中发生变化")

        for split_name, split_samples_ in splits.items():
            report["splits"][split_name] = {
                "count": len(split_samples_),
                "distribution": label_distribution(split_samples_),
            }

        REPORTS_DIR.mkdir(parents=True, exist_ok=True)
        REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"raw 数据清洗完成: train={len(splits['train'])}, dev={len(splits['dev'])}, test={len(splits['test'])}")
        print(f"报告: {REPORT_PATH}")
        return 0
    except Exception as error:
        print(f"[ERROR] {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
